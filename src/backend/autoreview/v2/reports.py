from __future__ import annotations
import csv
import io
import json
import math
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from .core import require, now, uid, atomic_write, dumps
from .telemetry import sanitize


def interval_ms(start, end):
    if not start or not end:
        return None
    return max(0, round((datetime.fromisoformat(end) - datetime.fromisoformat(start)).total_seconds() * 1000))


def statistics(service, task_id, group_ids=None):
    from .statistics_detail import info, wrong_answers
    store = service.store
    with store.connection() as db:
        groups = [g for g in store.list(db, 'groups', task_id=task_id, include_deleted=False)
                  if not group_ids or g['id'] in group_ids]
        ids = {g['id'] for g in groups}
        rows, versions, submissions, page_revisions = [], set(), [], {}
        references = {r['id']:r for r in store.list(db,'reference_versions',task_id=task_id)}
        metadata = {r['id']:r for r in store.list(db,'reference_metadata',task_id=task_id)}
        details = {}
        for sub in store.list(db, 'submissions', task_id=task_id, include_deleted=False):
            if sub['group_id'] not in ids:
                continue
            submissions.append(sub)
            if sub['state'] != 'reviewed':
                continue
            policy=store.get(db,'scoring_versions',sub.get('scoring_version_id') or '',optional=True)
            for page in store.list(db, 'pages', parent_id=sub['id']):
                if not page.get('reviewed'):
                    continue
                versions.add(page['reference_version_id'])
                page_revisions[page['id']]=page['revision']
                reference=references[page['reference_version_id']]
                answers={a['id']:a for a in reference['items']}
                for result in service.results(db, page):
                    key=(page['reference_version_id'],result['id'])
                    if key not in details:
                        details[key]=info(service,db,reference,answers[result['id']],result,page,
                            metadata.get(reference['id'],{}).get('name','答案 v'+str(reference.get('version',''))))
                    rows.append({'submission_id': sub['id'], 'paper_tag': sub['paper_tag'],
                        'group_id': sub['group_id'], 'page_id': page['id'], 'slot': page['slot'],
                        'reference_version_id': page['reference_version_id'], 'question_id': result['id'],
                        'number': result['number'], 'expected': result['english'], 'raw_english': result['raw_english'],
                        'status': result['status'], 'writing_status': result['writing_status'],
                        'row': result['row'], 'column': result['column'],
                        'label':details[key]['label'],'expected_display':details[key]['expected_display'],
                        'wrong_answers':wrong_answers(result,policy)})
        questions = {}
        for row in rows:
            key = (row['reference_version_id'], row['question_id'])
            q = questions.setdefault(key, {'reference_version_id': key[0], 'question_id': key[1],
                'number': row['number'], 'slot': row['slot'], 'expected': row['expected'],
                **details[key], 'known': 0, 'errors': 0, 'unknown': 0, 'writing_errors': 0,'wrong_counts':Counter()})
            q['unknown'] += row['status'] == 'unknown'
            q['known'] += row['status'] != 'unknown'
            q['errors'] += row['status'] in ('incorrect', 'missing')
            q['writing_errors'] += row['writing_status'] == 'confirmed_error'
            q['wrong_counts'].update(set(row['wrong_answers']))
        for q in questions.values():
            q['error_rate'] = q['errors'] / q['known'] if q['known'] else None
            q['wrong_answers']=[{'text':value,'count':count} for value,count in sorted(q.pop('wrong_counts').items(),key=lambda v:(-v[1],v[0]))]
        known = sum(q['known'] for q in questions.values())
        errors = sum(q['errors'] for q in questions.values())
        scores=[]
        for sub in submissions:
            if sub.get('score_snapshot_id'):
                score=store.get(db,'score_snapshots',sub['score_snapshot_id'])
                scores.append({**score,'paper_tag':sub['paper_tag']})
        score_groups={}
        for score in scores:
            key=score['scoring_version_id']
            score_groups.setdefault(key,{'scoring_version_id':key,'total_minor':score['total_minor'],'final_scores':[],'pending_count':0})
            if score['final_minor'] is None: score_groups[key]['pending_count']+=1
            else: score_groups[key]['final_scores'].append(score['final_minor'])
        for row in rows:
            score=next((v for v in scores if v['parent_id']==row['submission_id']),None)
            if score: row.update(scoring_version_id=score['scoring_version_id'],score_snapshot_id=score['id'],score=score['final_minor']/100 if score['final_minor'] is not None else None,total=score['total_minor']/100)
        return {'page_revisions':page_revisions,'scores':scores,'score_groups':list(score_groups.values()),'task_id': task_id, 'group_ids': sorted(ids), 'reference_versions': sorted(versions),
            'mixed_versions': len(versions) > 1, 'total_submissions': len(submissions),
            'reviewed_submissions': sum(s['state'] == 'reviewed' for s in submissions),
            'known': known, 'errors': errors, 'error_rate': errors / known if known else None,
            'unknown': sum(q['unknown'] for q in questions.values()),
            'questions': sorted(questions.values(), key=lambda q: (-(q['error_rate'] or 0), q['slot'], q['number'])),
            'rows': rows, 'generated_at': now()}


def metrics(service, task_id=None):
    store = service.store
    with store.connection() as db:
        attempts = store.list(db, 'attempts', task_id=task_id) + store.list(db, 'usage_ledger', task_id=task_id)
        runs = store.list(db, 'runs', task_id=task_id)
        sessions = store.list(db, 'activity_sessions', task_id=task_id)
        feedback = store.list(db, 'review_feedback', task_id=task_id)
        pages = store.list(db, 'pages', task_id=task_id)
        groups = store.list(db, 'groups', task_id=task_id, include_deleted=False)
        submissions = store.list(db, 'submissions', task_id=task_id, include_deleted=False)
        def summarize(selected):
            latencies = sorted(a['elapsed_ms'] for a in selected if a.get('elapsed_ms') is not None)
            usage = [a['usage'] for a in selected if a.get('usage') is not None]
            def total(*names):
                values = [next((u[n] for n in names if u.get(n) is not None), None) for u in usage]
                observed = [v for v in values if v is not None]
                return sum(observed) if observed else None
            return {'attempt_count': len(selected), 'failed_count': sum(a['state'] == 'failed' for a in selected),
                'interrupted_count': sum(a['state'] == 'interrupted' for a in selected),
                'known_inference_ms': sum(latencies) if latencies else None,
                'latency_observations': len(latencies), 'unknown_latency_count': len(selected) - len(latencies),
                'p50_ms': latencies[math.ceil(len(latencies) * .5) - 1] if latencies else None,
                'p95_ms': latencies[math.ceil(len(latencies) * .95) - 1] if latencies else None,
                'known_usage_count': len(usage), 'unknown_usage_count': len(selected) - len(usage),
                'input_tokens': total('inputTokens', 'input_tokens'), 'output_tokens': total('outputTokens', 'output_tokens'),
                'cached_input_tokens': total('cachedInputTokens', 'cached_input_tokens')}
        per_run = []
        for run in runs:
            selected = [a for a in attempts if a.get('run_id') == run['id']]
            per_run.append({**run, **summarize(selected),
                'wall_ms': interval_ms(run.get('submitted_at'), run.get('ended_at') or now())})
        by_prompt = {}
        for attempt in attempts:
            key = attempt.get('prompt_version_id')
            by_prompt.setdefault(key, []).append(attempt)
        feedback_attempt = {}
        for f in feedback:
            observation = store.get(db, 'observations', f['observation_id'], optional=True) if f.get('observation_id') else None
            if observation:
                feedback_attempt.setdefault(observation['parent_id'], set()).add(f['parent_id'])
        page_metrics = []
        for page in pages:
            related = [a for a in attempts if a.get('page_id') == page['id']]
            page_metrics.append({'page_id': page['id'], 'submission_id': page['parent_id'], 'group_id': page['group_id'],
                'slot': page['slot'], **summarize(related), 'import_to_processed_ms': interval_ms(page['created_at'], page.get('processed_at')),
                'first_wait_ms': interval_ms(page['created_at'], related[0].get('started_at')) if related else None,
                'reviewed': bool(page.get('reviewed'))})
        class_metrics = [{'group_id': g['id'], 'name': g['name'], **summarize([a for a in attempts if a.get('group_id') == g['id']]),
            'page_count': sum(p['group_id'] == g['id'] for p in pages),
            'reviewed_pages': sum(p['group_id'] == g['id'] and bool(p.get('reviewed')) for p in pages)} for g in groups]
        submission_metrics = [{'submission_id': s['id'], 'paper_tag': s['paper_tag'], 'group_id': s['group_id'],
            'state': s['state'], **summarize([a for a in attempts if a.get('page_id') in
                {p['id'] for p in pages if p['parent_id'] == s['id']}])} for s in submissions]
        return {'summary': summarize(attempts), 'attempts': sorted(attempts, key=lambda a: (a.get('started_at',a['created_at']),a['id']), reverse=True), 'runs': sorted(per_run, key=lambda r: (r.get('submitted_at',r['created_at']),r['id']), reverse=True), 'pages': page_metrics,
            'groups': class_metrics, 'submissions': submission_metrics,
            'human_active_ms': sum(s.get('active_ms', 0) for s in sessions if s.get('measurement') != 'paper_manual_timer'),
            'paper_manual_ms': sum((interval_ms(s['started_at'], now()) if s['state'] == 'running' else s.get('active_ms', 0))
                                   for s in sessions if s.get('measurement') == 'paper_manual_timer'),
            'running_timers': [s for s in sessions if s.get('measurement') == 'paper_manual_timer' and s['state'] == 'running'],
            'human_time_method': '估算：窗口可见且最近60秒有操作时的心跳累计；不代表节省时间',
            'reviewed_pages': sum(bool(p.get('reviewed')) for p in pages),
            'feedback_categories': dict(Counter(f['category'] for f in feedback)),
            'prompt_versions': [{'prompt_version_id': key, 'latest_at': max(a.get('started_at',a['created_at']) for a in value), **summarize(value),
                'teacher_corrected_page_count': len(set().union(*(feedback_attempt.get(a['id'], set()) for a in value)))}
                for key, value in sorted(by_prompt.items(), key=lambda pair: max(a.get('started_at',a['created_at']) for a in pair[1]), reverse=True)],
            'quota_note': '账户共享窗口仅作观测；不由Token换算单图额度，不把缺失值记为0。'}


def annotate(service, page_id, frozen=None, expected_revision=None):
    with service.store.connection() as db:
        page = service.store.get(db, 'pages', page_id)
        require(expected_revision is None or page['revision'] == expected_revision, 'RESULT_CHANGED', '批注结果已更新，请刷新图片。', 409)
        version = service.store.get(db, 'image_versions', frozen['image_version_id'] if frozen else page['image_version_id'])
        results = frozen['results'] if frozen else service.results(db, page)
    image = Image.open(service.asset(version['asset_id'])).convert('RGB')
    draw = ImageDraw.Draw(image)
    width, height = image.size
    line = max(2, width // 350)
    from .annotations import marks
    for r in results:
        for mark in r.get('annotation_marks',marks(r)):
            box=mark['box']
            color = '#e63946' if mark['status'] in ('incorrect', 'missing') else '#d28c16'
            bounds = [box['x1'] * width, box['y1'] * height, box['x2'] * width, box['y2'] * height]
            draw.ellipse(bounds, outline=color, width=line)
            draw.text((bounds[0], max(0, bounds[1] - 16)), mark['label'], fill=color)
    return image


def export_report(service, task_id, data):
    require(data.get('format') in ('csv', 'xlsx', 'json', 'pdf', 'images'), 'EXPORT_FORMAT', '不支持该格式。')
    stats = statistics(service, task_id, data.get('group_ids'))
    require(stats['reviewed_submissions'], 'NOT_REVIEWED', '所选班级还没有完整复核的作业。')
    store = service.store
    with store.connection(write=True) as db:
        pages = []
        for page_id in dict.fromkeys(r['page_id'] for r in stats['rows']):
            page = store.get(db, 'pages', page_id)
            require(page['revision']==stats['page_revisions'][page_id], 'RESULT_CHANGED', '复核结果已更新，请重新导出。', 409)
            pages.append({'page_id': page_id, 'image_version_id': page['image_version_id'],
                'review_revision_id': page['history'][page['history_cursor']] if page.get('history_cursor', -1) >= 0 else None,
                'results': service.results(db, page)})
        snapshot = store.put(db, 'report_snapshots', {'parent_id': task_id, 'task_id': task_id,
            'statistics': stats, 'pages': pages, 'name': data.get('name') or '复核结果'})
    directory = Path(service.settings.data['export_directory'])
    directory.mkdir(parents=True, exist_ok=True)
    identifier = uid()
    kind = data['format']
    suffix = 'zip' if kind == 'images' else kind
    destination = directory / f'AutoReview-{snapshot["id"][:8]}.{suffix}'
    columns = ['paper_tag', 'group_id', 'slot', 'number', 'expected', 'raw_english', 'status', 'writing_status', 'reference_version_id', 'scoring_version_id', 'score_snapshot_id', 'score', 'total']
    if kind == 'json':
        atomic_write(destination, dumps(snapshot).encode('utf-8'))
    elif kind == 'csv':
        stream = io.StringIO(newline='')
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction='ignore')
        writer.writeheader()
        for row in stats['rows']:
            # Prevent formulas when the teacher opens untrusted student text in Excel.
            writer.writerow({k: ("'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@', '\t', '\r')) else v) for k, v in row.items()})
        atomic_write(destination, stream.getvalue().encode('utf-8-sig'))
    elif kind == 'xlsx':
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = '复核结果'
        sheet.append(columns)
        for row in stats['rows']:
            sheet.append([row.get(k) for k in columns])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        summary = book.create_sheet('错题统计')
        keys = ['slot', 'number', 'expected', 'known', 'errors', 'unknown', 'error_rate']
        summary.append(keys)
        for q in stats['questions']:
            summary.append([q[k] for k in keys])
            # Reference text and question labels can also start with '='.
            # Keep both sheets literal, just like student transcription above.
            for cell in summary[summary.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        stream = io.BytesIO()
        book.save(stream)
        atomic_write(destination, stream.getvalue())
    else:
        images = [annotate(service, p['page_id'], p) for p in pages]
        stream = io.BytesIO()
        if kind == 'pdf':
            images[0].save(stream, 'PDF', save_all=True, append_images=images[1:])
        else:
            with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
                for page, image in zip(pages, images):
                    output = io.BytesIO()
                    image.save(output, 'PNG')
                    archive.writestr(page['page_id'] + '.png', output.getvalue())
                archive.writestr('results.json', dumps(snapshot))
        atomic_write(destination, stream.getvalue())
        for image in images:
            image.close()
    with store.connection(write=True) as db:
        record = store.put(db, 'exports', {'id': identifier, 'parent_id': snapshot['id'], 'task_id': task_id,
            'format': kind, 'output_path': str(destination), 'state': 'completed'})
        store.event(db, 'report.exported', record, format=kind, snapshot_id=snapshot['id'])
    return record


def export_diagnostics(service, task_id=None):
    result = metrics(service, task_id)
    # Whitelist metadata. No image, reference answers, transcription, prompts, identity, or credential files.
    attempts = [{k: a.get(k) for k in ('id', 'run_id', 'purpose', 'state', 'elapsed_ms', 'usage', 'error_code',
                    'prompt_version_id', 'prompt_hash', 'schema_hash', 'preprocessing_version', 'grading_rule_version',
                    'actual_model', 'actual_service_tier')} for a in result['attempts']]
    with service.store.connection() as db:
        events = [dict(row) for row in db.execute('SELECT * FROM events WHERE (? IS NULL OR task_id=?) ORDER BY event_id', (task_id, task_id))]
    payload = sanitize({'format_version': 1, 'created_at': now(), 'summary': result['summary'],
        'attempts': attempts, 'diagnostics': service.diagnostics.rows(2000), 'events': events,
        'prompt_versions': result['prompt_versions'], 'feedback_categories': result['feedback_categories'],
        'limitations': '教师修正属于迭代线索，尚非评测真值；不自动修改提示词或重新调用模型。'})
    destination = Path(service.settings.data['export_directory']) / f'AutoReview-diagnostics-{uid()[:8]}.json'
    atomic_write(destination, dumps(payload).encode('utf-8'))
    return {'path': str(destination)}
